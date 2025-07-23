from functools import partial, reduce 
from pathlib import Path
from typing import Union

from openpyxl import load_workbook
from openpyxl.utils import exceptions as ee
import pandas as pd
from toolz import functoolz as fz

from .helper_0 import shortcut_target, str_camel_to_snake


class partial2(partial):   
    """An improved version of partial that uses Ellipsis (...) as a placeholder."""
    def __init__(self, func, *args, kwargs_first=None, **kwargs): 
        if ((kwargs_first is None)
            and (any(x is ... for x in args))
            and (any(v is ... for x in kwargs.values()))): 
            warn("Positional arguments with ... are evaluated before keyword ones."
                "\nUse kwargs_first=True to control behavior.")
            kwargs_first = False
        self.kwargs_first = kwargs_first

    def __call__(self, *args, **keywords):
        # notation is tricky: 
        # η for iterator, λ for lambdas, 0 intermediate vars, 1 final vars.  
        η_args = iter(args)
        λ_ellipsis = lambda x: next(η_args) if x is ... else x
        if self.kwargs_first: 
            keywords_0 = {k: λ_ellipsis(v) for k, v in self.keywords.items()}
            args_1 = tuple(λ_ellipsis(arg) for arg in self.args) + tuple(η_args)
        else: 
            args_0 = (λ_ellipsis(arg) for arg in self.args)
            keywords_0 = {k: λ_ellipsis(v) for k, v in self.keywords.items()}
            args_1 = tuple(args_0) + tuple(η_args)
        keywords_1 = {**keywords_0, **keywords}        
        return self.func(*args_1, **keywords_1)


def star(func):
    # Unpacks star operator:  star(func)(args) := func(*args) 
    return lambda args: func(*args)


def thread(val, *forms):
    # Unify pytoolz.thread_(first|last) with Ellipsis ...
    eval_ff = (lambda vv, ff: 
        partial2(*ff)(vv) if isinstance(ff, tuple) else ff(vv))
    return reduce(eval_ff, forms, val)
    

def read_excel_table(
    a_file: Union[str, Path], a_sheet: str, a_table: str=None, **kwargs
    ) -> pd.DataFrame:
    """(file, sheet, table) returns pandas.DataFrame of values."""
    import pandas as pd
    a_file = Path(a_file)
    
    if a_table is None:
        str_camel_2_snake = fz.compose_left(str_plus, ϱ('to_snake'))
        a_table = str_camel_2_snake(a_sheet)

    try:
        a_wb = load_workbook(a_file, data_only=True)
    except (ee.InvalidFileException, PermissionError) as e:
        b_file = shortcut_target(a_file)
        a_wb = load_workbook(b_file, data_only=True)

    a_ws  = a_wb[a_sheet]
    a_tbl = a_ws.tables[a_table]
    rows_ls = [[ cell.value for cell in row ] for row in a_ws[a_tbl.ref]]
    tbl_df  = pd.DataFrame(data=rows_ls[1:],
            index=None, columns=rows_ls[0], **kwargs)
    return tbl_df