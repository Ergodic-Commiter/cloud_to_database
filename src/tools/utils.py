from operator import attrgetter as ɑ
from pathlib import Path, WindowsPath
from platform import system
import re
from typing import Union

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries
from openpyxl.worksheet.table import Table as XLTable
# pylint:disable=import-outside-toplevel
# pylint:disable=broad-exception-raised
# pylint:disable=import-error
# pylint:disable=super-init-not-called
# pylint:disable=too-many-arguments


class OpenTable(XLTable): 
    """Lightweight wrapper around openpyxl.Table"""
    __module__ = 'src.tools'

    def __init__(self, wb_path:Union[str,Path], ws_name:str, tb_name:str, *, 
            data_only=True, read_only=False):
        self.wb_path = wb_path
        self.ws_name = ws_name
        self.tb_name = tb_name
        self.wb_args = dict(data_only=data_only, read_only=read_only)
        (wb, _) = self.check_workbook(wb_path, **self.wb_args)
        ws = wb[ws_name]
        self.workbook = wb
        self._table = ws.tables[tb_name]

    def __getattr__(self, name): 
        return getattr(self._table, name)
    
    def __repr__(self): 
        repr_str = ("<OpenTable wraps Table(name={0}, ref={1}) at (path={2}, sheet={3})>"
            .format(*ɑ('tb_name', '_table.ref', 'wb_path', 'ws_name')(self)))
        return repr_str
    
    def get_dataframe(self, **kwargs):
        import pandas as pd
        a_ws = self.get_sheet() 
        λ_name = lambda nn: nn.strip().replace(' ', '_')
        rows_ls = [map(ɑ('value'), row) for row in a_ws[self._table.ref]]
        the_df = pd.DataFrame(data=rows_ls[1:], 
            columns=map(λ_name, rows_ls[0]), index=None, **kwargs)
        return the_df

    @property
    def boundaries(self): 
        return range_boundaries(self._table.ref)

    @staticmethod
    def check_workbook(a_path:Union[str,Path], *, data_only, read_only): 
        wb_args = dict(data_only=data_only, read_only=read_only)
        try: 
            a_wb = load_workbook(a_path, **wb_args)
        except PermissionError: 
            a_path = shortcut_target(a_path)
            a_wb = load_workbook(a_path, **wb_args)
        return (a_wb, a_path)

    def get_sheet(self): 
        return self.workbook[self.ws_name]    


def shortcut_target(a_file:Union[str,Path], **kwargs):
    a_file = Path(a_file)
    if a_file.is_symlink(): 
        return a_file.resolve()
    if system() == 'Windows' and a_file.suffix == '.lnk': 
        return _windows_shortcut(a_file, kwargs.get('file_ext'))
    if system() == 'Darwin': 
        return _mac_alias(a_file)
    raise ValueError("Nor a recognized shortcut or symlink.")

def _mac_alias(a_file:Path): 
    from mac_alias import read_alias   
    with a_file.open('rb') as _f: 
        return read_alias(_f).path

def _windows_shortcut(a_file:Path, file_ext:str=None):
    if file_ext is None:
        if isinstance(a_file, WindowsPath):
            file_ext = re.findall(r"\.([A-Za-z]{3,4})\.lnk", a_file.name)[0]
        else:
            raise Exception("Couldn't determine file extension.")
    file_regex = fr'(C:\\.*\.{file_ext})'
    with open(a_file, 'r', encoding='ISO-8859-1') as _f:
        a_path = re.findall(file_regex, _f.read(), flags=re.DOTALL)
    if len(a_path) != 1:
        raise Exception('Not unique or no shortcut targets found in link.')
    return a_path[0]

def _windows_shortcut_2(a_file:Path): 
    from pylnk3 import parse
    with a_file.open('rb') as f: 
        return Path(parse(f).path)

def str_camel_to_snake(a_str:str): 
    raise NotImplementedError("Function 'str_camel_to_snake' is not implemented.")