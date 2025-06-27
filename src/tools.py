from operator import attrgetter, methodcaller as ϱ
from pathlib import WindowsPath, Path
import re
from typing import Union
from warnings import warn

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
import pandas as pd
from toolz import compose_left



def read_excel_table(
    a_file: Union[str, Path], a_sheet: str, a_table: str=None, **kwargs
    ) -> pd.DataFrame:
    """(file, sheet, table) returns pandas.DataFrame of values."""
    a_file = Path(a_file)
    
    if a_table is None:
        str_camel_2_snake = compose_left(str_plus, ϱ('to_snake'))
        a_table = str_camel_2_snake(a_sheet)

    try:
        a_wb = load_workbook(a_file, data_only=True)
    except InvalidFileException:
        b_file = _shortcut_target(a_file)
        a_wb = load_workbook(b_file, data_only=True)

    a_ws  = a_wb[a_sheet]
    a_tbl = a_ws.tables[a_table]
    rows_ls = [[ cell.value for cell in row ] for row in a_ws[a_tbl.ref]]
    tbl_df  = pd.DataFrame(data=rows_ls[1:],
            index=None, columns=rows_ls[0], **kwargs)
    return tbl_df


def _shortcut_target(a_file: Union[str, Path], file_ext:str=None):
    if isinstance(a_file, str):
        a_file = Path(a_file)

    if file_ext is None:
        if isinstance(a_file, WindowsPath):
            file_ext = re.findall(r"\.([A-Za-z]{3,4})\.lnk", a_file.name)[0]
        else:
            raise Exception("Couldn't determine file extension.")

    file_regex = fr'(C:\\.*\.{file_ext})'
    with open(a_file, 'r', encoding='ISO-8859-1') as _f:
        a_path = re.findall(file_regex, _f.read(), flags=re.DOTALL)

    if len(a_path) != 1:
        raise Exception('Not unique or no shortcut targets found in link.')
    return a_path[0]
