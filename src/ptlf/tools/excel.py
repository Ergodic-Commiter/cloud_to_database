from operator import attrgetter as ɑ
from pathlib import Path, WindowsPath
from platform import system
import re

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries
from openpyxl.worksheet.table import Table as XLTable
# pylint:disable=super-init-not-called
# pylint:disable=too-many-arguments
# pylint:disable=import-outside-top-level


class OpenTable(XLTable): 
    """Lightweight wrapper around openpyxl.Table"""
    __module__ = 'ptlf.tools'

    def __init__(self, wb_path:Path, ws_name:str, tb_name:str, *, 
            data_only=True, read_only=False):
        self.wb_path = Path(wb_path)
        self.ws_name = ws_name
        self.tb_name = tb_name
        self.wb_args = dict(data_only=data_only, read_only=read_only)
        wb = self.check_workbook(wb_path, **self.wb_args)
        self.workbook = wb
        self._table = wb[ws_name].tables[tb_name]

    def __getattr__(self, name): 
        return getattr(self._table, name)
    
    def __repr__(self):
        attrs = ɑ('tb_name', '_table.ref', 'wb_path', 'ws_name')(self)
        repr_str = ("<OpenTable wraps Table(name={0}, ref={1}) at (path={2}, sheet={3})>"
            .format(*attrs))
        return repr_str
    
    def get_dataframe(self, **kwargs):
        import pandas as pd
        a_ws = self.get_sheet() 
        λ_name = lambda nn: nn.strip().replace(' ', '_')
        rows_ls = [map(ɑ('value'), row) for row in a_ws[self._table.ref]]
        the_df = pd.DataFrame(data=rows_ls[1:], 
            columns=map(λ_name, rows_ls[0]), index=None, **kwargs)
        return the_df

    @property
    def boundaries(self): 
        return range_boundaries(self._table.ref)

    @staticmethod
    def check_workbook(a_path:Path, *, data_only, read_only): 
        a_path = Path(a_path)
        wb_args = dict(data_only=data_only, read_only=read_only)
        try: 
            return load_workbook(a_path, **wb_args)
        except PermissionError: 
            a_path = ShortcutPath(a_path).get_target()
            return load_workbook(a_path, **wb_args)

    def get_sheet(self): 
        return self.workbook[self.ws_name]    


class ShortcutPath(Path): 
    def __init__(self, obj:Path): 
        self.obj = Path(obj)
    
    def get_target(self): 
        if self.obj.is_symlink():
            return self.obj.resolve()
        system_ = system()
        if system_ == 'Windows': 
            return self._windows_shortcut()
        if system_ == 'Darwin': 
            return self._mac_alias()
        raise ValueError(f"Can't get shortcut target in system {system_}")

    def _mac_alias(self):
        # Tramposo porque Mac tiene tanto Alias como Shortcuts y no son lo mismo. 
        from mac_alias import read_alias  # pylint: disable=no-name-in-module
        with self.obj.open('rb') as _f: 
            return read_alias(_f).path

    def _windows_shortcut(self):
        if isinstance(self.obj, WindowsPath):
            file_ext = re.findall(r"\.([A-Za-z]{3,4})\.lnk", self.obj.name)[0]
        else:
            raise ValueError("Couldn't determine file extension.")
        file_regex = fr'(C:\\.*\.{file_ext})'
        with self.obj.open('r', encoding='ISO-8859-1') as _f:
            a_path = re.findall(file_regex, _f.read(), flags=re.DOTALL)
        if len(a_path) != 1:
            raise ValueError('Not unique or no shortcut targets found in link.')
        return a_path[0]

    def _windows_shortcut_2(self): 
        from pylnk3 import parse
        with self.obj.open('rb') as _f: 
            return Path(parse(_f).path)

